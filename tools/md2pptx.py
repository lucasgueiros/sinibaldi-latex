#!/usr/bin/env python3
"""Converte <stem>.slides.md (Markdown revisado pelo LLM) em <stem>.pptx.

    python3 tools/md2pptx.py vol2/02_secao_x/01_cap_y/01_cap_y.slides.md
    # -> vol2/02_secao_x/01_cap_y/01_cap_y.pptx

Convenção de entrada (simples de propósito; a fonte versionada é o .md):
  - slides separados por uma linha só com `---`
  - `# Título`  -> slide de seção (só título)
  - `## Título` -> título do slide
  - `> ...`     -> citação do livro (vai como parágrafo, marcada)
  - `- ...` / `* ...` -> tópico
  - demais linhas não vazias -> parágrafo de apoio

Requer python-pptx (ver requirements.txt): .venv/bin/pip install python-pptx
"""
import argparse
import os
import re
import sys

try:
    from pptx import Presentation
    from pptx.util import Pt
except ModuleNotFoundError:
    sys.exit("falta python-pptx:  .venv/bin/pip install python-pptx")

INLINE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`")


def limpa(linha: str) -> str:
    linha = INLINE.sub(lambda m: next(g for g in m.groups() if g is not None), linha)
    return linha.strip()


def blocos(md: str):
    atual = []
    for ln in md.splitlines():
        if ln.strip() == "---":
            if atual:
                yield atual
                atual = []
        else:
            atual.append(ln)
    if atual:
        yield atual


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("md")
    ap.add_argument("-o", "--out", help="padrão: <md sem .slides.md>.pptx")
    a = ap.parse_args()

    md = open(a.md, encoding="utf-8").read()
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)          # tira o comentário-instrução

    prs = Presentation()
    layout_secao = prs.slide_layouts[5]      # só título
    layout_conteudo = prs.slide_layouts[1]   # título + corpo

    for linhas in blocos(md):
        titulo, corpo, secao = "", [], False
        for ln in linhas:
            s = ln.strip()
            if not s:
                continue
            if s.startswith("## "):
                titulo = limpa(s[3:])
            elif s.startswith("# "):
                titulo, secao = limpa(s[2:]), True
            elif s.startswith(("- ", "* ")):
                corpo.append(("bullet", limpa(s[2:])))
            elif s.startswith(">"):
                t = limpa(s[1:].strip())
                if t:
                    corpo.append(("quote", t))
            else:
                corpo.append(("par", limpa(s)))

        if not titulo and not corpo:
            continue
        if secao and not corpo:
            sl = prs.slides.add_slide(layout_secao)
            sl.shapes.title.text = titulo
            continue

        sl = prs.slides.add_slide(layout_conteudo)
        sl.shapes.title.text = titulo or " "
        tf = sl.placeholders[1].text_frame
        tf.word_wrap = True
        for i, (tipo, txt) in enumerate(corpo):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = ("“" + txt + "”") if tipo == "quote" else txt
            p.level = 1 if tipo == "bullet" else 0
            if tipo == "quote":
                for r in p.runs:
                    r.font.italic = True
                    r.font.size = Pt(16)

    saida = a.out or re.sub(r"(\.slides)?\.md$", "", a.md) + ".pptx"
    prs.save(saida)
    print(f"escrito: {saida}  ({len(prs.slides._sldIdLst)} slide(s))")


if __name__ == "__main__":
    main()
